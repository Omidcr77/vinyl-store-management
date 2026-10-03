from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image, ImageDraw
import shutil

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
ASSETS.mkdir(exist_ok=True)
SOURCES = {
    "intake": ROOT / "test-results/delivery-review.png",
    "sale": ROOT / "test-results/multi-sale.png",
    "account": ROOT / "test-results/customer-statement.png",
    "suppliers": ROOT / "test-results/supplier-mobile.png",
    "receipt": ROOT / "test-results/receipt-preview.png",
}
for key, source in SOURCES.items():
    if not source.exists():
        raise FileNotFoundError(source)
    shutil.copy2(source, ASSETS / f"{key}.png")
for weight in (400, 700):
    font_path = ROOT / f"node_modules/@fontsource/vazirmatn/files/vazirmatn-arabic-{weight}-normal.woff2"
    if font_path.exists():
        shutil.copy2(font_path, ASSETS / f"vazirmatn-{weight}.woff2")

W, H = 13.333333, 7.5
BG, INK, GREEN, MINT, GOLD = "F5F3EC", "18352D", "1F5948", "DDE9DF", "D9A84E"
MUTED, WHITE, LINE = "66766F", "FFFFFF", "DCE2D9"
FONT = "Tahoma"

slides = [
    {"kind":"cover", "eyebrow":"راه ساده‌تر برای مدیریت دکان شما", "title":"هر رول،\nهر فروش،\nهمه در یک‌جا.", "sub":"سیستم مدیریت دکان قالین و وینیل", "tag":"موجودی • فروشات • حساب مشتریان • تهیه‌کنندگان", "image":"intake"},
    {"eyebrow":"01  /  موجودی", "title":"از موجودی دکان‌تان دقیق باخبر باشید.", "lead":"هر رول را جداگانه ثبت کنید و طول باقی‌ماندهٔ آن را بدانید.", "bullets":[("رسیدن اجناس را یک‌جا ثبت کنید","تهیه‌کننده، رنگ، اندازه و مصارف هر محموله را وارد کنید."),("هر رول را جداگانه نگه دارید","برای هر رول یک ثبت جدا و طول باقی‌ماندهٔ آن را داشته باشید."),("کمبود موجودی را زود ببینید","موجودی آمادهٔ فروش و رول‌های کم‌موجود را از داشبورد بررسی کنید.")], "image":"intake", "caption":"بررسی محموله • هر رول به‌صورت جداگانه ثبت می‌شود", "pill":"موجودی دقیق‌تر، غافلگیری کمتر"},
    {"eyebrow":"02  /  فروش", "title":"فروش را در یک بل به‌سادگی ثبت کنید.", "lead":"مشتری و جنس‌های انتخاب‌شده را وارد کنید و پرداخت را یک‌بار ثبت نمایید.", "bullets":[("یک یا چند رول اضافه کنید","طول و نرخ هر جنس را جداگانه در همان بل بنویسید."),("مجموع بل را خودکار ببینید","قیمت جنس‌ها، پرداخت و باقی‌داری در یک‌جا محاسبه می‌شود."),("موجودی را هم‌زمان به‌روز کنید","طول فروخته‌شده از رول انتخاب‌شده کم می‌شود.")], "image":"sale", "caption":"صفحهٔ فروش چندجنسه • یک مشتری و چند جنس در یک بل", "pill":"فروش منظم و سریع‌تر"},
    {"eyebrow":"03  /  حساب مشتریان", "title":"حساب هر مشتری را روشن و منظم نگه دارید.", "lead":"خریدها، پرداخت‌ها و باقی‌داری مشتری را در یک سابقهٔ کامل ببینید.", "bullets":[("سابقهٔ حساب را یک‌جا ببینید","خریدها و رسیدهای بعدی در صورت‌حساب مشتری نمایش داده می‌شود."),("پرداخت‌های بعدی را ثبت کنید","با هر رسید، باقی‌داری حساب به‌روز می‌شود."),("بل و رسید آماده کنید","بل، رسید و صورت‌حساب را PDF کنید؛ اندازهٔ چاپ آن A4 است.")], "image":"account", "caption":"صورت‌حساب مشتری • خریدها، رسیدها و باقی‌داری فعلی", "pill":"حساب‌وکتاب روشن‌تر"},
    {"eyebrow":"04  /  تهیه‌کنندگان", "title":"مصارف خرید و حساب تهیه‌کنندگان را یک‌جا ببینید.", "lead":"جنس‌های وارده را با حساب تهیه‌کننده و مصارف خرید پیوند دهید.", "bullets":[("مصارف خرید را ثبت کنید","قیمت خرید فی متر و مصارف اضافی را هنگام ورود جنس بنویسید."),("حساب تهیه‌کننده را دنبال کنید","پرداخت، پیش‌پرداخت، قرض و صورت‌حساب تهیه‌کننده را ثبت کنید."),("گزارش را با درنظرداشت مصرف بخوانید","سود ناخالص از قیمت خرید ثبت‌شده حساب می‌شود؛ هزینهٔ نامعلوم مشخص می‌ماند.")], "image":"suppliers", "caption":"حساب تهیه‌کننده • جزئیات و فعالیت در موبایل", "pill":"آگاهی بهتر از مصارف"},
    {"eyebrow":"05  /  کار روزمره", "title":"برای کار روزمرهٔ یک دکان ساخته شده.", "lead":"ابزارهای ساده برای ثبت سفارش، مدیریت موجودی و حساب‌وکتاب دکان.", "bullets":[("صفحه‌ها و اسناد به زبان دری","رابط کاربری راست‌به‌چپ و اسناد چاپی به زبان دری هستند."),("در کمپیوتر و موبایل کار می‌کند","صفحه‌ها برای انجام کارهای اصلی دکان در هر دو آماده‌اند."),("دسترسی بر اساس وظیفه","نقش‌های مدیر سیستم، مدیر و کارمند دسترسی هر شخص را تنظیم می‌کنند.")], "image":"receipt", "caption":"رسید چاپی • سند A4 برای شریک‌ساختن با مشتری", "pill":"آمادهٔ استفادهٔ روزانه"},
    {"kind":"close", "eyebrow":"یک گام عملی برای دکان شما", "title":"کارهای دکان‌تان را منظم‌تر پیش ببرید.", "sub":"موجودی، فروشات، حساب مشتریان و مصارف تهیه‌کنندگان را در یک جریان کاری مدیریت کنید.", "cta":"بیایید کار روزمرهٔ دکان‌تان را باهم مرور کنیم.", "contact":"نام شما  •  شماره تماس / واتس‌اپ  •  ایمیل", "image":"sale"},
]

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
blank = prs.slide_layouts[6]

def rect(slide,x,y,w,h,fill,radius=False,line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill)
    shape.line.fill.background() if not line else None
    if line:
        shape.line.color.rgb=rgb(line); shape.line.width=Pt(1)
    if radius:
        try: shape.adjustments[0]=0.1
        except Exception: pass
    return shape

def text(slide,s,x,y,w,h,size=16,color=INK,bold=False,font=FONT,align=None,vertical=MSO_ANCHOR.MIDDLE):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0); tf.vertical_anchor=vertical
    for i,line in enumerate(s.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.alignment=align if align is not None else PP_ALIGN.RIGHT
        p._p.get_or_add_pPr().set("rtl", "1")
        p.space_after=Pt(2)
    return box

def photo(slide,key,x,y,w,h):
    p=ASSETS/f"{key}.png"
    bg=rect(slide,x,y,w,h,WHITE,True,LINE)
    with Image.open(p) as im:
        ratio=im.width/im.height
    avail_w,avail_h=w-0.20,h-0.20
    pic_w=min(avail_w,avail_h*ratio); pic_h=pic_w/ratio
    if pic_h>avail_h: pic_h=avail_h; pic_w=pic_h*ratio
    pic=slide.shapes.add_picture(str(p), Inches(x+(w-pic_w)/2), Inches(y+(h-pic_h)/2), width=Inches(pic_w), height=Inches(pic_h))
    return pic

def add_footer(slide,n):
    text(slide,"سیستم مدیریت دکان قالین و وینیل",0.62,7.16,5.5,0.17,8,MUTED,True)
    text(slide,f"{n:02d}  /  07",11.82,7.13,0.85,0.21,8,MUTED,True,align=PP_ALIGN.RIGHT)

for idx,d in enumerate(slides,1):
    s=prs.slides.add_slide(blank)
    bg=s.background.fill; bg.solid(); bg.fore_color.rgb=rgb(INK if d.get('kind') in ('cover','close') else BG)
    if d.get("kind")=="cover":
        rect(s,0.72,0.78,0.11,0.55,GOLD,True)
        text(s,d["eyebrow"].upper(),1.00,0.78,5.6,0.38,11,"D7E6DC",True)
        text(s,d["title"],0.77,1.57,5.4,3.05,35,WHITE,True,vertical=MSO_ANCHOR.TOP)
        text(s,d["sub"],0.81,4.93,5.2,0.45,18,"DDE9DF",True)
        text(s,d["tag"],0.82,5.49,5.2,0.42,12,"B7CCBE")
        # Screenshot panel with a product frame.
        rect(s,6.36,1.36,6.28,4.55,"FFFFFF",True)
        photo(s,d["image"],6.52,1.48,5.96,4.31)
        text(s,"تصویر از محیط برنامه",0.82,6.84,2.6,0.19,8,"A8C1B1",True)
        text(s,"01 / 07",11.80,6.84,0.75,0.19,8,"A8C1B1",True,align=PP_ALIGN.RIGHT)
    elif d.get("kind")=="close":
        text(s,d["eyebrow"].upper(),0.82,0.83,6.4,0.35,11,"B7CCBE",True)
        text(s,d["title"],0.82,1.48,6.0,1.68,31,WHITE,True,vertical=MSO_ANCHOR.TOP)
        text(s,d["sub"],0.84,3.28,5.4,0.9,17,"D7E6DC",False,vertical=MSO_ANCHOR.TOP)
        rect(s,0.82,4.45,5.5,0.75,GREEN,True)
        text(s,d["cta"],1.04,4.55,5.1,0.52,14,WHITE,True)
        text(s,d["contact"],0.84,5.49,5.6,0.45,10,"C1D1C7")
        photo(s,d["image"],6.65,0.98,5.84,5.68)
        text(s,"معلومات تماس خود را این‌جا بنویسید",0.84,6.86,3.8,0.19,8,"A8C1B1",True)
        text(s,"07 / 07",11.80,6.84,0.75,0.19,8,"A8C1B1",True,align=PP_ALIGN.RIGHT)
    else:
        text(s,d["eyebrow"],0.70,0.45,5.5,0.27,9,GREEN,True)
        text(s,d["title"],0.70,0.91,5.35,0.78,25,INK,True,vertical=MSO_ANCHOR.TOP)
        text(s,d["lead"],0.72,1.78,5.12,0.76,12,MUTED,False,vertical=MSO_ANCHOR.TOP)
        for j,(head,body) in enumerate(d["bullets"]):
            yy=2.72+j*1.12
            rect(s,0.72,yy,0.48,0.48,MINT,True)
            text(s,f"{j+1:02d}",0.72,yy+0.01,0.48,0.43,10,GREEN,True,align=PP_ALIGN.CENTER)
            text(s,head,1.37,yy-0.02,4.35,0.31,13,INK,True,vertical=MSO_ANCHOR.TOP)
            text(s,body,1.37,yy+0.34,4.38,0.57,10,MUTED,False,vertical=MSO_ANCHOR.TOP)
        rect(s,0.72,6.21,4.95,0.47,GREEN,True)
        text(s,d["pill"],0.94,6.25,4.48,0.35,10,WHITE,True)
        # screenshot column
        rect(s,6.12,1.10,6.47,4.98,WHITE,True,LINE)
        photo(s,d["image"],6.30,1.25,6.11,4.54)
        text(s,d["caption"],6.34,6.20,5.94,0.23,8,MUTED,False)
        add_footer(s,idx)

prs.core_properties.title="معرفی سیستم مدیریت دکان قالین و وینیل"
prs.core_properties.subject="Product overview, customer benefits and daily use"
prs.core_properties.author=""
prs.save(OUT/"flooring-vinyl-store-manager.pptx")

# Matching print-ready HTML version for a PDF handout.
def esc(v):
    import html
    return html.escape(v).replace("\n","<br>")
pages=[]
for idx,d in enumerate(slides,1):
    if d.get('kind')=='cover':
        pages.append(f'''<section class="slide cover"><div class="covercopy"><small>{esc(d['eyebrow']).upper()}</small><h1>{esc(d['title'])}</h1><h2>{esc(d['sub'])}</h2><p>{esc(d['tag'])}</p><small class="foot">تصویر از محیط برنامه</small></div><div class="heroimg"><img src="assets/{d['image']}.png"></div><span class="pageno">01 / 07</span></section>''')
    elif d.get('kind')=='close':
        pages.append(f'''<section class="slide cover close"><div class="covercopy"><small>{esc(d['eyebrow']).upper()}</small><h1>{esc(d['title'])}</h1><h2 class="closelead">{esc(d['sub'])}</h2><div class="cta">{esc(d['cta'])}</div><p>{esc(d['contact'])}</p><small class="foot">معلومات تماس خود را این‌جا بنویسید</small></div><div class="heroimg"><img src="assets/{d['image']}.png"></div><span class="pageno">07 / 07</span></section>''')
    else:
        bullets=''.join(f'<div class="bullet"><b>{j:02d}</b><div><strong>{esc(h)}</strong><p>{esc(body)}</p></div></div>' for j,(h,body) in enumerate(d['bullets'],1))
        pages.append(f'''<section class="slide content"><div class="copy"><small>{esc(d['eyebrow'])}</small><h1>{esc(d['title'])}</h1><p class="lead">{esc(d['lead'])}</p>{bullets}<div class="pill">{esc(d['pill'])}</div></div><div class="shot"><img src="assets/{d['image']}.png"><p>{esc(d['caption'])}</p></div><span class="brand">سیستم مدیریت دکان قالین و وینیل</span><span class="pageno">{idx:02d} / 07</span></section>''')
html_doc='''<!doctype html><html lang="fa-AF" dir="rtl"><head><meta charset="utf-8"><style>
@font-face{font-family:Vazirmatn;src:url('assets/vazirmatn-400.woff2') format('woff2');font-weight:400}@font-face{font-family:Vazirmatn;src:url('assets/vazirmatn-700.woff2') format('woff2');font-weight:700}
@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Vazirmatn,Arial,'DejaVu Sans',sans-serif;color:#18352D}.slide{width:13.333in;height:7.5in;position:relative;overflow:hidden;page-break-after:always;background:#F5F3EC;padding:.57in}.cover{direction:ltr;background:#18352D;color:#fff;display:flex;padding:.76in;gap:.5in}.covercopy{direction:rtl;text-align:right;width:5.35in;position:relative;padding-top:.12in}.cover small{font-size:10pt;letter-spacing:0;color:#B7CCBE;font-weight:bold}.cover h1{font-size:32pt;line-height:1.22;margin:.48in 0 .3in;white-space:pre-line;color:#fff}.cover h2{font-size:16pt;color:#DDE9DF;margin:0 0 .18in}.cover p{font-size:11pt;color:#B7CCBE}.heroimg{margin-left:auto;width:6.05in;height:4.55in;margin-top:.62in;background:#fff;border-radius:16px;padding:.13in;display:flex;align-items:center;justify-content:center}.heroimg img{width:100%;height:100%;object-fit:contain}.foot{position:absolute;bottom:0;left:0}.pageno{position:absolute;right:.62in;bottom:.28in;font-size:9pt;color:#66766F;direction:ltr}.cover .pageno{color:#A8C1B1}.content{direction:ltr;display:flex;gap:.42in;padding:.55in .68in}.copy{direction:rtl;text-align:right;width:5.03in;position:relative}.copy>small{font-size:9pt;font-weight:bold;color:#1F5948;letter-spacing:0}.copy h1{font-size:24pt;line-height:1.25;margin:.20in 0 .16in}.lead{font-size:11pt;line-height:1.5;color:#66766F;margin:0 0 .18in}.bullet{direction:rtl;text-align:right;display:flex;gap:.2in;margin:.16in 0;min-height:.82in}.bullet b{background:#DDE9DF;color:#1F5948;border-radius:9px;width:.46in;height:.46in;text-align:center;padding-top:.13in;font-size:10pt;flex:none}.bullet strong{font-size:12pt}.bullet p{margin:.06in 0;color:#66766F;font-size:9pt;line-height:1.4}.pill{position:absolute;bottom:.44in;left:0;background:#1F5948;color:white;border-radius:8px;padding:.13in .2in;width:4.95in;font-weight:bold;font-size:10pt;text-align:right}.shot{direction:rtl;text-align:right;margin-left:auto;margin-top:.55in;width:6.42in;height:4.98in;border-radius:14px;border:1px solid #DCE2D9;background:white;padding:.14in}.shot img{width:100%;height:4.5in;object-fit:contain}.shot p{color:#66766F;font-size:8pt;margin:.05in .04in}.brand{position:absolute;bottom:.19in;left:.68in;color:#66766F;font-weight:bold;font-size:8pt}.close h1{font-size:29pt;margin:.48in 0 .22in}.closelead{font-size:15pt;line-height:1.5!important;font-weight:normal}.cta{margin-top:.42in;background:#1F5948;border-radius:9px;padding:.22in;color:white;font-size:13pt;font-weight:bold}.close .covercopy>p{font-size:10pt;margin-top:.25in}.close .foot{bottom:0}.close .heroimg{height:4.55in;margin-top:.8in}
</style></head><body>''' + ''.join(pages) + '</body></html>'
(OUT/"flooring-vinyl-store-manager.html").write_text(html_doc,encoding='utf-8')
print("Created PowerPoint and matching HTML deck in", OUT)
